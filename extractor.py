"""
extractor.py
------------
Extracts raw text from a PDF (PyMuPDF) or a DOCX (python-docx) file.
Preserves reading order block-by-block, page-by-page.
"""

from __future__ import annotations

import fitz  # PyMuPDF
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def extract_text_from_pdf(pdf_path: str) -> str:
    """
    Extract all text from a PDF file in reading order.

    Args:
        pdf_path: Absolute path to the PDF file.

    Returns:
        A single string containing all extracted text,
        with pages separated by a form-feed character.

    Raises:
        ValueError: If the PDF cannot be opened or is encrypted.
    """
    try:
        doc = fitz.open(pdf_path)
    except Exception as exc:
        raise ValueError(f"Cannot open PDF '{pdf_path}': {exc}") from exc

    if doc.is_encrypted:
        raise ValueError("Encrypted PDFs are not supported.")

    pages_text: list[str] = []

    for page_num in range(len(doc)):
        page = doc.load_page(page_num)

        # "blocks" preserves spatial reading order (top -> bottom, left -> right)
        blocks = page.get_text("blocks", sort=True)

        page_lines: list[str] = []
        for block in blocks:
            # block = (x0, y0, x1, y1, text, block_no, block_type)
            # block_type 0 = text, 1 = image
            if block[6] == 0:
                raw = block[4].strip()
                if raw:
                    page_lines.append(raw)

        pages_text.append("\n".join(page_lines))

    doc.close()

    # Join pages; form-feed lets the parser split on pages if needed
    return "\f".join(pages_text)


def extract_metadata(pdf_path: str) -> dict:
    """
    Extract PDF metadata (title, author, etc.) as a supplemental hint.

    Returns:
        dict with keys: title, author, subject, creator, producer
    """
    try:
        doc = fitz.open(pdf_path)
        meta = doc.metadata or {}
        doc.close()
        return {
            "title": meta.get("title", ""),
            "author": meta.get("author", ""),
            "subject": meta.get("subject", ""),
            "creator": meta.get("creator", ""),
            "producer": meta.get("producer", ""),
        }
    except Exception:
        return {}


# ---------------------------------------------------------------------------
# Photo / image extraction
# ---------------------------------------------------------------------------
# Headshots ride along inside resume files as embedded images. We pull out the
# LARGEST web-renderable image (photos are always the biggest; icons, bullets
# and divider lines are tiny) and save it to disk, so the JSON can reference
# it by URL instead of base64.

_PHOTO_MIME = {
    "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
    "webp": "image/webp", "gif": "image/gif", "bmp": "image/bmp",
}
_MIN_PHOTO_DIM   = 60      # px — anything smaller is an icon/bullet/line
_MIN_PHOTO_BYTES = 1200    # bytes — skip 1-px spacers etc.


def _photo_mime(ext: str) -> str | None:
    return _PHOTO_MIME.get(ext.lower())


def _photo_ext_allowed(ext: str) -> bool:
    return ext.lower() in _PHOTO_MIME


def extract_photo_from_pdf(pdf_path: str, out_base: str) -> str | None:
    """
    Extract the largest embedded image (the headshot) from a PDF.

    Args:
        pdf_path: Path to the source PDF.
        out_base: Output path WITHOUT extension; the correct extension is
                  appended based on the image's real format.

    Returns:
        The path actually written (e.g. '<out_base>.jpg') or None if the PDF
        contains no usable photo. Never raises.
    """
    try:
        doc = fitz.open(pdf_path)
    except Exception:
        return None

    best: tuple[int, int] | None = None   # (pixel_area, xref)
    try:
        for page in doc:
            for img in page.get_images(full=True):
                xref, width, height = img[0], img[2], img[3]
                if width < _MIN_PHOTO_DIM or height < _MIN_PHOTO_DIM:
                    continue
                area = width * height
                if best is None or area > best[0]:
                    best = (area, xref)
        if best is None:
            return None

        xref = best[1]
        info = doc.extract_image(xref)
        ext = (info.get("ext") or "").lower()
        if ext == "jpeg":
            ext = "jpg"

        if _photo_ext_allowed(ext):
            out = f"{out_base}.{ext}"
            with open(out, "wb") as fh:
                fh.write(info["image"])
            return out

        # Unusual container (jp2/jpx/jxr/...): decode and re-save as PNG.
        pix = fitz.Pixmap(doc, xref)
        if pix.colorspace is not None and pix.colorspace.n > 3:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        out = f"{out_base}.png"
        pix.save(out)
        return out
    except Exception:
        return None
    finally:
        doc.close()


def extract_photo_from_docx(docx_path: str, out_base: str) -> str | None:
    """
    Extract the largest embedded image from a Word document.

    Same contract as extract_photo_from_pdf. Never raises.
    """
    try:
        doc = Document(docx_path)
    except Exception:
        return None

    best: tuple[int, bytes, str] | None = None   # (size, blob, ext)
    try:
        for part in doc.part.package.parts:
            ctype = getattr(part, "content_type", "") or ""
            if not ctype.startswith("image/"):
                continue
            ext = ctype.split("/", 1)[1].split("+")[0].lower()
            if ext == "jpeg":
                ext = "jpg"
            if not _photo_ext_allowed(ext):
                continue          # skip EMF/WMF/SVG — browsers can't <img> them
            blob = part.blob
            if len(blob) < _MIN_PHOTO_BYTES:
                continue
            if best is None or len(blob) > best[0]:
                best = (len(blob), blob, ext)
        if best is None:
            return None
        out = f"{out_base}.{best[2]}"
        with open(out, "wb") as fh:
            fh.write(best[1])
        return out
    except Exception:
        return None


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------

_W_T   = qn('w:t')
_W_TAB = qn('w:tab')
_W_BR  = qn('w:br')
_W_CR  = qn('w:cr')
_W_P   = qn('w:p')
_W_TBL = qn('w:tbl')


def _docx_paragraph_text(paragraph) -> str:
    """Rebuild paragraph text including tabs and soft line breaks."""
    parts = []
    for node in paragraph._p.iter():
        if node.tag == _W_T:
            parts.append(node.text or '')
        elif node.tag == _W_TAB:
            parts.append('\t')
        elif node.tag in (_W_BR, _W_CR):
            parts.append('\n')
    return ''.join(parts)


def _docx_block_text(el, parent) -> str:
    """Text of one top-level body element (paragraph or table), in document order."""
    if el.tag == _W_P:
        return _docx_paragraph_text(Paragraph(el, parent))

    if el.tag == _W_TBL:
        rows = []
        for row in Table(el, parent).rows:
            cells = [_Cell(c, row).text.replace('\n', ' ').strip() for c in row._tr.tc_lst]
            rows.append('\t'.join(c for c in cells if c))
        return '\n'.join(r for r in rows if r)

    return None


def extract_text_from_docx(docx_path: str) -> str:
    """
    Extract all text from a .docx in reading order.

    Walks the document body in order so paragraphs and tables interleave the
    way they appear on the page (a naive `doc.paragraphs` loop silently drops
    every table). Tabs are rendered as ' | ' so parser._split_role_company()
    can split role vs company. Exact-duplicate lines are collapsed.

    Raises:
        ValueError: If the DOCX cannot be opened.
    """
    try:
        doc = Document(docx_path)
    except Exception as exc:
        raise ValueError(f"Cannot open DOCX '{docx_path}': {exc}") from exc

    lines: list[str] = []
    for el in doc.element.body.iterchildren():
        chunk = _docx_block_text(el, doc)
        if chunk:
            for ln in chunk.replace('\t', ' | ').split('\n'):
                ln = ln.strip()
                if ln:
                    lines.append(ln)

    seen, deduped = set(), []
    for ln in lines:
        if ln not in seen:
            seen.add(ln)
            deduped.append(ln)

    return '\n'.join(deduped)


def extract_metadata_from_docx(docx_path: str) -> dict:
    """Same dict shape as extract_metadata(), sourced from Word core properties."""
    try:
        props = Document(docx_path).core_properties
        return {
            "title":    props.title or "",
            "author":   props.author or "",
            "subject":  props.subject or "",
            "creator":  props.last_modified_by or "",
            "producer": "Microsoft Word",
        }
    except Exception:
        return {}
